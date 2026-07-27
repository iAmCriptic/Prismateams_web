"""Document format preference for new office-like files (OOXML vs OpenDocument)."""

from __future__ import annotations

import zipfile
from typing import Literal

from app.models.settings import SystemSettings

SETTING_DOCUMENT_FORMAT = 'files_document_format'
FORMAT_OFFICE = 'office'
FORMAT_OPENDOCUMENT = 'opendocument'
DocumentFormat = Literal['office', 'opendocument']

_OFFICE_TYPES = {
    'document': 'docx',
    'spreadsheet': 'xlsx',
    'presentation': 'pptx',
}
_ODF_TYPES = {
    'document': 'odt',
    'spreadsheet': 'ods',
    'presentation': 'odp',
}

_MIME_TYPES = {
    'docx': 'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
    'xlsx': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    'pptx': 'application/vnd.openxmlformats-officedocument.presentationml.presentation',
    'odt': 'application/vnd.oasis.opendocument.text',
    'ods': 'application/vnd.oasis.opendocument.spreadsheet',
    'odp': 'application/vnd.oasis.opendocument.presentation',
}

_ODF_CONTENT = {
    'odt': '''<?xml version="1.0" encoding="UTF-8"?>
<office:document-content xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0"
 xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0" office:version="1.2">
 <office:body>
  <office:text>
   <text:p/>
  </office:text>
 </office:body>
</office:document-content>
''',
    'ods': '''<?xml version="1.0" encoding="UTF-8"?>
<office:document-content xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0"
 xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0"
 xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0" office:version="1.2">
 <office:body>
  <office:spreadsheet>
   <table:table table:name="Sheet1">
    <table:table-column/>
    <table:table-row>
     <table:table-cell><text:p/></table:table-cell>
    </table:table-row>
   </table:table>
  </office:spreadsheet>
 </office:body>
</office:document-content>
''',
    'odp': '''<?xml version="1.0" encoding="UTF-8"?>
<office:document-content xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0"
 xmlns:draw="urn:oasis:names:tc:opendocument:xmlns:drawing:1.0"
 xmlns:presentation="urn:oasis:names:tc:opendocument:xmlns:presentation:1.0"
 xmlns:svg="urn:oasis:names:tc:opendocument:xmlns:svg-compatible:1.0"
 xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0" office:version="1.2">
 <office:body>
  <office:presentation>
   <draw:page draw:name="page1">
    <draw:frame svg:width="10cm" svg:height="2cm" svg:x="2cm" svg:y="2cm">
     <draw:text-box><text:p/></draw:text-box>
    </draw:frame>
   </draw:page>
  </office:presentation>
 </office:body>
</office:document-content>
''',
}

_ODF_STYLES = '''<?xml version="1.0" encoding="UTF-8"?>
<office:document-styles xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" office:version="1.2">
 <office:styles/>
</office:document-styles>
'''

_ODF_META = '''<?xml version="1.0" encoding="UTF-8"?>
<office:document-meta xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" office:version="1.2">
 <office:meta/>
</office:document-meta>
'''


def get_document_format() -> DocumentFormat:
    setting = SystemSettings.query.filter_by(key=SETTING_DOCUMENT_FORMAT).first()
    value = (setting.value if setting else FORMAT_OFFICE) or FORMAT_OFFICE
    value = str(value).strip().lower()
    if value in (FORMAT_OPENDOCUMENT, 'odf', 'opendoc'):
        return FORMAT_OPENDOCUMENT
    return FORMAT_OFFICE


def get_create_type_map(fmt: DocumentFormat | None = None) -> dict[str, str]:
    """Return {document, spreadsheet, presentation} → file extension."""
    fmt = fmt or get_document_format()
    return dict(_ODF_TYPES if fmt == FORMAT_OPENDOCUMENT else _OFFICE_TYPES)


def get_allowed_create_types(fmt: DocumentFormat | None = None) -> set[str]:
    return set(get_create_type_map(fmt).values())


def mime_for_type(file_type: str) -> str:
    return _MIME_TYPES[file_type]


def create_empty_document(filepath: str, file_type: str) -> str:
    """Create an empty document on disk. Returns MIME type."""
    if file_type not in _MIME_TYPES:
        raise ValueError(f'Unsupported file type: {file_type}')

    if file_type in ('docx', 'xlsx', 'pptx'):
        if file_type == 'docx':
            from docx import Document
            Document().save(filepath)
        elif file_type == 'xlsx':
            from openpyxl import Workbook
            Workbook().save(filepath)
        else:
            from pptx import Presentation
            Presentation().save(filepath)
        return _MIME_TYPES[file_type]

    _write_odf(filepath, file_type)
    return _MIME_TYPES[file_type]


def _write_odf(filepath: str, file_type: str) -> None:
    mimetype = _MIME_TYPES[file_type]
    content_xml = _ODF_CONTENT[file_type]
    manifest = (
        '<?xml version="1.0" encoding="UTF-8"?>\n'
        '<manifest:manifest xmlns:manifest="urn:oasis:names:tc:opendocument:xmlns:manifest:1.0" '
        'manifest:version="1.2">\n'
        f' <manifest:file-entry manifest:full-path="/" manifest:version="1.2" '
        f'manifest:media-type="{mimetype}"/>\n'
        ' <manifest:file-entry manifest:full-path="content.xml" manifest:media-type="text/xml"/>\n'
        ' <manifest:file-entry manifest:full-path="styles.xml" manifest:media-type="text/xml"/>\n'
        ' <manifest:file-entry manifest:full-path="meta.xml" manifest:media-type="text/xml"/>\n'
        '</manifest:manifest>\n'
    )
    with zipfile.ZipFile(filepath, 'w') as zf:
        # ODF requires uncompressed mimetype as first entry
        zf.writestr('mimetype', mimetype, compress_type=zipfile.ZIP_STORED)
        zf.writestr('META-INF/manifest.xml', manifest)
        zf.writestr('content.xml', content_xml)
        zf.writestr('styles.xml', _ODF_STYLES)
        zf.writestr('meta.xml', _ODF_META)
